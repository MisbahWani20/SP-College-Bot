
import sqlite3
from pathlib import Path
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parent
EXCEL_PATH = (
    PROJECT_DIR
    / "data"
    / "official_data"
    / "knowledge"
    / "sp_college_departments.xlsx"
)
DATABASE_PATH = PROJECT_DIR / "chatbot" / "knowledge.db"

if not EXCEL_PATH.exists():
    raise FileNotFoundError(f"Excel file not found: {EXCEL_PATH}")

workbook = load_workbook(EXCEL_PATH, data_only=True)
sheet = workbook.active

headers = [
    str(cell.value).strip() if cell.value is not None else ""
    for cell in sheet[1]
]

required_columns = [
    "College",
    "Department",
    "Programmes Offered",
    "Level",
    "Category",
    "Facilities",
    "Hostel Available",
    "Hostel Fee",
    "HOD",
    "Scholarships Available / Type",
    "Overview",
    "Source",
    "Official URL",
    "HOD Contact",
]

missing = [column for column in required_columns if column not in headers]

if missing:
    raise ValueError(
        f"Missing Excel columns: {missing}\n"
        f"Found columns: {headers}"
    )

rows = []

for values in sheet.iter_rows(min_row=2, values_only=True):
    record = dict(zip(headers, values))

    if not record.get("Department"):
        continue

    rows.append(tuple(
        str(record.get(column) or "").strip()
        for column in required_columns
    ))

if not rows:
    raise ValueError("No department records found in the Excel file.")

with sqlite3.connect(DATABASE_PATH) as conn:
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sp_college_departments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            college TEXT,
            department TEXT NOT NULL,
            programmes_offered TEXT,
            level TEXT,
            category TEXT,
            facilities TEXT,
            hostel_available TEXT,
            hostel_fee TEXT,
            hod TEXT,
            scholarships TEXT,
            overview TEXT,
            source TEXT,
            official_url TEXT,
            hod_contact TEXT
        )
    """)

    # Replace previous imported records to prevent duplicates.
    conn.execute("DELETE FROM sp_college_departments")

    conn.executemany("""
        INSERT INTO sp_college_departments (
            college,
            department,
            programmes_offered,
            level,
            category,
            facilities,
            hostel_available,
            hostel_fee,
            hod,
            scholarships,
            overview,
            source,
            official_url,
            hod_contact
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, rows)
    conn.commit()
print(f"{len(rows)} department records imported successfully.")
